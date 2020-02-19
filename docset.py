# coding:utf-8
# import os
# import win32com
# from win32com.client import Dispatch, constants
from docx import Document
import shelve, string
from copy import deepcopy
from config_menu import frame_style
from config_set_menu import item_format


def testfun(file):
    docf = Document(file)
    tab = docf.tables[0]
    index = 1
    item_dict = {}
    item_list = []
    while True:
        try:
            itemName = tab.cell(index, 1).text
            itemName = itemName.strip()
            itemType = tab.cell(index, 2).text
            iType = chkType(itemType)
            if iType == -1:
                return (-1, index)  # 第index行格式错误

            itemDisp = tab.cell(index, 3).text
            iDisp = chkDisp(itemDisp)
            if iDisp == -1:
                return (-1, index)  # 第index行格式错误

            itemScale = tab.cell(index, 4).text
            try:
                itemScale = eval(itemScale)
            except:
                return (-1, index)
            # iName=itemName.strip()
            item_list.append(itemName)
            newitem = item_format(name=itemName, No=index, dataStyle=iType, disp=iDisp, scale=itemScale)
            item_dict[itemName] = newitem
            index += 1
            # print('%-12s%-12s%-12s' % (itemName,itemType,itemDisp))

        except:
            break
    newFrame = frame_style('New Frame')
    newFrame.frame_datCfg_list = deepcopy(item_list)
    newFrame.frame_datCfg_dict = deepcopy(item_dict)
    return (1, newFrame)


def chkType(type):
    x = type.upper().strip()
    sty_dict = {'UINT32': 0, 'INT32': 1, 'FLOAT32': 2, 'FLOAT40': 3,
                'FLOAT64': 4, 'UINT16': 5, 'INT16': 6, 'UINT8': 7, 'INT8': 8}
    if x in sty_dict.keys():
        return sty_dict[x]
    else:
        return -1


def chkDisp(disp):
    x = disp.upper().strip()
    dis_dict = {'DEC': 0, 'HEX': 1}
    if x in dis_dict.keys():
        return dis_dict[x]
    else:
        return -1


if __name__ == '__main__':
    testfun('test.docx')